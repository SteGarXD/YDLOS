#!/usr/bin/env python3
"""Generate DataLens editions comparison xlsx for YDLOS."""
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    raise SystemExit("pip install openpyxl")

ROWS = [
    ("Возможность", "Cloud", "OSS official 2.9", "On-prem vendor", "YDLOS", "Комментарий"),
    ("Интерактивные дашборды", "✓", "✓", "✓", "✓", "vendor"),
    ("Конструктор чартов", "✓", "✓", "✓", "✓", "~37 типов HC=0"),
    ("Рассылки по расписанию", "2026", "—", "2026", "—", "отдельный сервис"),
    ("Отчёты PDF / презентации", "✓", "—", "✓", "частично", "PDF D3; не report builder"),
    ("Стилизация интерфейса", "✓", "—", "✓", "частично", "logo, SERVICE_NAME, theme scss"),
    ("Editor (JS) + API connector", "✓", "—", "✓", "—", "не в YDLOS"),
    ("Встраивание непубличных", "✓", "—", "2026", "частично", "share/embed overlay"),
    ("Яндекс.Карты", "✓", "—", "✓*", "частично", "YANDEX_MAP env"),
    ("Публичные чарты/дашборды", "✓", "—", "—", "частично", "shared links"),
    ("Авторизация (роли)", "✓", "—*", "✓", "✓", "*OSS+official auth; YDLOS full"),
    ("Аутентификация", "Yandex ID", "Local", "SSO/Local", "✓", "auth+OIDC"),
    ("Свои кастомные роли", "✓", "—", "✓", "частично", "legacy /admin/roles"),
    ("Создание пользователей", "✓", "частично", "✓", "✓", "/settings или legacy"),
    ("Проекты (изоляция)", "✓", "—", "✓", "частично", "legacy pd_projects или workbooks"),
    ("Группы пользователей", "✓", "—", "✓", "—", ""),
    ("Usage Analytics", "✓", "—", "✓", "—", "письмо Yandex"),
    ("Фоновый экспорт CSV", "✓", "частично", "✓", "частично", ""),
    ("Работа с файлами", "✓", "—", "✓", "—", ""),
    ("API объектов (US)", "✓", "✓", "✓", "✓", "US 1.39"),
    ("Контроль публикации", "✓", "частично", "✓", "частично", "US permissions"),
    ("Excel дашборда", "✓", "частично", "✓", "✓", "EXPORT_DASH_EXCEL"),
    ("Export workbook", "✓", "✓", "✓", "✓", "env"),
    ("QL __user_id / __embed", "✓", "—", "✓", "✓", "overlay"),
    ("Связанные объекты", "частично", "—", "частично", "✓", "overlay"),
    ("Share view-only", "✓", "частично", "✓", "✓", "overlay"),
    ("AI агент", "✓", "—", "2026", "—", ""),
    ("Обновление с official", "—", "—", "vendor", "✓", "behind=0, publish-dist"),
]

def main() -> None:
    root = Path(__file__).resolve().parents[1]
    out = root / "docs" / "dev" / "DataLens-editions-YDLOS.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Editions"
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")
    for r, row in enumerate(ROWS, 1):
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if r == 1:
                cell.fill = header_fill
                cell.font = header_font
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["F"].width = 36
    wb.save(out)
    print(f"Wrote {out}")

if __name__ == "__main__":
    main()
